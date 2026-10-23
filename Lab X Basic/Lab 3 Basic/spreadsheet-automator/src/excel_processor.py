# spreadsheet-automator/src/excel_processor.py
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

class ExcelProcessor:
    """
    คลาสสำหรับรวบรวมฟังก์ชันการทำงานกับ Excel โดยใช้ openpyxl
    """
    def __init__(self):
        pass

    def create_workbook(self, filename="new_workbook.xlsx"):
        """สร้าง Excel workbook ใหม่"""
        wb = openpyxl.Workbook()
        if 'Sheet' in wb.sheetnames:
            del wb['Sheet']
        print(f"สร้าง Workbook ใหม่ '{filename}' เรียบร้อยแล้ว")
        return wb

    def load_workbook(self, filename):
        """โหลด Excel workbook ที่มีอยู่แล้ว"""
        try:
            wb = openpyxl.load_workbook(filename)
            print(f"โหลด Workbook '{filename}' สำเร็จ")
            return wb
        except FileNotFoundError:
            print(f"ข้อผิดพลาด: ไม่พบ Workbook '{filename}'")
            return None
        except Exception as e:
            print(f"เกิดข้อผิดพลาดในการโหลด Workbook '{filename}': {e}")
            return None

    def save_workbook(self, workbook, filename):
        """บันทึก Workbook ลงในไฟล์ที่กำหนด"""
        try:
            workbook.save(filename)
            print(f"บันทึก Workbook ลงใน '{filename}' สำเร็จ")
            return True
        except Exception as e:
            print(f"เกิดข้อผิดพลาดในการบันทึก Workbook ลงใน '{filename}': {e}")
            return False

    def process_sales_data(self, input_filename, output_filename):
        """
        โหลดข้อมูลยอดขาย คำนวณยอดรวม และเขียนลงในไฟล์ Excel ใหม่พร้อมจัดรูปแบบ
        """
        input_wb = self.load_workbook(input_filename)
        if not input_wb:
            return False

        input_sheet = input_wb.active

        output_wb = self.create_workbook()
        output_sheet = output_wb.create_sheet("Sales Report", 0)

        # กำหนดรูปแบบ (Styles)
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4CAF50", end_color="4CAF50", fill_type="solid") # สีเขียว
        total_row_font = Font(bold=True, color="000000")
        total_row_fill = PatternFill(start_color="FFEB3B", end_color="FFEB3B", fill_type="solid") # สีเหลือง
        thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
        currency_format = '#,##0.00'

        # เขียนส่วนหัวลงในแผ่นงานใหม่
        headers = ["Product Name", "Quantity", "Unit Price", "Total Price"]
        for col_num, header in enumerate(headers, 1):
            cell = output_sheet.cell(row=1, column=col_num, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = thin_border
            cell.alignment = Alignment(horizontal='center', vertical='center')

        grand_total = 0
        output_row = 2

        # วนลูปอ่านข้อมูลทีละแถว (ข้ามแถวหัวตาราง)
        for row_index, row in enumerate(input_sheet.iter_rows(min_row=2, values_only=True), 2):
            if not any(row): # ข้ามแถวว่าง
                continue

            product_name = row[0] if len(row) > 0 else None
            quantity = row[1] if len(row) > 1 else 0
            unit_price = row[2] if len(row) > 2 else 0.0

            try:
                quantity = int(quantity)
                unit_price = float(unit_price)
            except (ValueError, TypeError):
                print(f"คำเตือน: ข้ามแถวที่ {row_index} เนื่องจากรูปแบบตัวเลขไม่ถูกต้อง: {row}")
                continue

            total_price = quantity * unit_price
            grand_total += total_price

            # เขียนข้อมูลลงในแผ่นงานใหม่
            output_sheet.cell(row=output_row, column=1, value=product_name).border = thin_border
            output_sheet.cell(row=output_row, column=2, value=quantity).border = thin_border
            output_sheet.cell(row=output_row, column=3, value=unit_price).border = thin_border
            
            total_price_cell = output_sheet.cell(row=output_row, column=4, value=total_price)
            total_price_cell.number_format = currency_format
            total_price_cell.border = thin_border

            output_row += 1

        # เขียนแถวยอดขายรวมทั้งหมด (Grand Total)
        output_sheet.merge_cells(start_row=output_row, start_column=1, end_row=output_row, end_column=3)
        grand_total_label_cell = output_sheet.cell(row=output_row, column=1, value="Grand Total:")
        grand_total_label_cell.font = total_row_font
        grand_total_label_cell.fill = total_row_fill
        grand_total_label_cell.alignment = Alignment(horizontal='right', vertical='center')

        grand_total_value_cell = output_sheet.cell(row=output_row, column=4, value=grand_total)
        grand_total_value_cell.font = total_row_font
        grand_total_value_cell.fill = total_row_fill
        grand_total_value_cell.number_format = currency_format
        grand_total_value_cell.border = thin_border

        # ปรับความกว้างคอลัมน์อัตโนมัติ
        for col_num in range(1, len(headers) + 1):
            max_length = 0
            column_letter = get_column_letter(col_num)
            for cell in output_sheet[column_letter]:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except:
                    pass
            adjusted_width = (max_length + 2)
            if adjusted_width > 0:
                output_sheet.column_dimensions[column_letter].width = adjusted_width

        return self.save_workbook(output_wb, output_filename)