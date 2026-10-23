# document-automator/src/word_processor.py
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os
from utils import setup_logging

logger = setup_logging(__name__)

class WordProcessor:
    """จัดการการสร้างและแก้ไขเอกสาร Word ด้วย python-docx"""
    def __init__(self):
        logger.info("WordProcessor initialized.")

    def create_document(self):
        """สร้างเอกสาร Word เปล่าขึ้นมาใหม่"""
        document = docx.Document()
        logger.info("New Word document created.")
        return document

    def add_heading(self, document, text, level=1):
        """เพิ่มหัวข้อ (Heading) ลงในเอกสาร"""
        document.add_heading(text, level=level)
        logger.info(f"Added heading (Level {level}): '{text}'")

    def add_paragraph(self, document, text, style=None, bold=False, italic=False, font_size=None, align='left'):
        """เพิ่มย่อหน้าพร้อมกำหนดรูปแบบ (Formatting)"""
        paragraph = document.add_paragraph(text, style=style)
        if paragraph.runs:
            run = paragraph.runs[0]
            if bold:
                run.bold = True
            if italic:
                run.italic = True
            if font_size:
                run.font.size = Pt(font_size)
                
        if align == 'center':
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == 'right':
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        elif align == 'justify':
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        else:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            
        logger.info(f"Added paragraph: '{text[:50]}...'")

    def add_table(self, document, data, include_header=True, header_bold=True, col_widths=None):
        """เพิ่มตารางลงในเอกสาร"""
        if not data:
            logger.warning("No data provided to add table.")
            return
            
        rows = len(data)
        cols = len(data[0]) if data else 0
        if rows == 0 or cols == 0:
            logger.warning("Table data is empty or malformed.")
            return
            
        table = document.add_table(rows=rows, cols=cols)
        table.style = 'Table Grid'
        
        if col_widths and len(col_widths) == cols:
            for i, width in enumerate(col_widths):
                table.columns[i].width = Inches(width)
                
        for r_idx, row_data in enumerate(data):
            row_obj = table.rows[r_idx]
            for c_idx, cell_data in enumerate(row_data):
                cell = row_obj.cells[c_idx]
                cell.text = str(cell_data)
                if include_header and r_idx == 0 and header_bold:
                    if cell.paragraphs[0].runs:
                        cell.paragraphs[0].runs[0].bold = True
                        
        logger.info(f"Added table with {rows} rows and {cols} columns.")

    def save_document(self, document, output_path):
        """บันทึกเอกสาร Word ไปยังพาทที่กำหนด"""
        try:
            document.save(output_path)
            logger.info(f"Word document saved to '{output_path}'.")
            return True
        except Exception as e:
            logger.error(f"Error saving Word document to '{output_path}': {e}")
            return False